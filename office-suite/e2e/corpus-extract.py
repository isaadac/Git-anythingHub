"""What other programs see in an Office file: the reference the corpus test compares Isadek Office with.

    python3 e2e/corpus-extract.py FILE...    prints one JSON object per file (one per line)

Uses python-docx, openpyxl and python-pptx, which read the files independently of Isadek Office.
  .docx  {"words": [...]}  the words of the body (paragraphs and tables, in order)
  .xlsx  {"sheets": {name: {"A1": value, ...}}, "formulas": n}  values as last calculated
  .pptx  {"slides": n, "words": [...]}  the words on the slides
An unreadable file gives {"error": "..."}.
"""
import datetime
import json
import re
import sys
import warnings

warnings.filterwarnings("ignore")
W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"


def words(text):
    return [w for w in re.split(r"\s+", text) if w]


def docx(path):
    import docx as d
    doc = d.Document(path)
    out = []
    # Paragraphs (body and tables) in order; text in w:t, not deleted text.
    for p in doc.element.body.iter(W + "p"):
        # Paragraphs inside text boxes are counted by their own w:p, not twice.
        out += words("".join(_text(t) for t in p.iter() if t.tag in RUN_TEXT and _own(t, p)))
    return {"words": out}


# Text and what separates words in a run: tabs and line breaks.
RUN_TEXT = {W + "t", W + "tab", W + "br", W + "cr"}


def _text(t):
    return (t.text or "") if t.tag == W + "t" else " "


def _own(t, p):
    x = t.getparent()
    while x is not None and x is not p:
        if x.tag == W + "p":
            return False
        x = x.getparent()
    return True


def value(v):
    if isinstance(v, (datetime.datetime, datetime.date, datetime.time)):
        return {"date": v.isoformat()}
    if isinstance(v, datetime.timedelta):
        return v.total_seconds() / 86400
    if isinstance(v, (bool, int, float, str)):
        return v
    return str(v)


def xlsx(path):
    import openpyxl
    wb = openpyxl.load_workbook(path, data_only=True)
    sheets = {}
    for ws in wb.worksheets:
        cells = {}
        for row in ws.iter_rows():
            for c in row:
                if c.value is not None and c.value != "":
                    cells[c.coordinate] = value(c.value)
        sheets[ws.title] = cells
    formulas = 0
    try:
        wf = openpyxl.load_workbook(path, data_only=False)
        for ws in wf.worksheets:
            for row in ws.iter_rows():
                for c in row:
                    if isinstance(c.value, str) and c.value.startswith("="):
                        formulas += 1
    except Exception:
        pass
    return {"sheets": sheets, "formulas": formulas}


def pptx(path):
    from pptx import Presentation
    prs = Presentation(path)
    out = []
    for slide in prs.slides:
        for shape in slide.shapes:
            out += _shape_words(shape)
    return {"slides": len(prs.slides), "words": out}


def _shape_words(shape):
    out = []
    if shape.shape_type == 6:  # group
        for s in shape.shapes:
            out += _shape_words(s)
        return out
    if getattr(shape, "has_table", False) and shape.has_table:
        for row in shape.table.rows:
            for cell in row.cells:
                out += words(cell.text)
        return out
    if getattr(shape, "has_text_frame", False) and shape.has_text_frame:
        for p in shape.text_frame.paragraphs:
            out += words("".join(r.text for r in p.runs) + "".join(f.text or "" for f in p._p.iter(A + "fld") if f.text))
    return out


def main(paths):
    for path in paths:
        ext = path.rsplit(".", 1)[-1].lower()
        try:
            r = {"docx": docx, "xlsx": xlsx, "pptx": pptx}[ext](path)
        except Exception as e:  # what other programs would report as a broken file
            r = {"error": type(e).__name__ + ": " + str(e)[:300]}
        r["file"] = path
        print(json.dumps(r, ensure_ascii=False))


if __name__ == "__main__":
    main(sys.argv[1:])
